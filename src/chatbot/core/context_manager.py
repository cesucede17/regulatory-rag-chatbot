"""Team context manager for Tambora slash commands."""

from __future__ import annotations

import re
from dataclasses import dataclass
from pathlib import Path

# src/chatbot/core/context_manager.py -> parents[3] es la raiz del proyecto
CONTEXTS_DIR = Path(__file__).resolve().parents[3] / "contexto"

BASE_SYSTEM_PROMPT = (
    "Eres Tambora, un asistente experto en legislación española y el BOE.\n"
    "Responde ÚNICAMENTE con información presente en el CONTEXTO proporcionado.\n\n"
    "Reglas estrictas:\n"
    "- Si el contexto está vacío o no contiene información relevante, responde exactamente: "
    '"No he podido recuperar información para esa consulta."\n'
    "- No inventes títulos, referencias, reales decretos, leyes ni URLs.\n"
    "- Reproduce títulos y URLs exactamente como aparecen en el contexto, sin modificarlos.\n"
    "- Usa formato Markdown: listas con viñetas, negrita para títulos, cursiva para fechas.\n"
    "- Sé conciso pero completo: incluye todos los resultados relevantes del contexto.\n"
    "- Nunca respondas con información fuera del contexto proporcionado.\n\n"
    "Instrucciones de formato de respuesta:\n"
    "- Lista los resultados relevantes con viñetas (-).\n"
    "- Para cada resultado usa la estructura: **Título** [Rango] — Departamento — URL (fecha).\n"
    "- Si el contexto proviene de PDFs: cita el nombre del documento y la página entre corchetes.\n"
    "- Si el usuario pide una sección concreta (ej: solo Ministerio de Hacienda), filtra y muestra "
    "solo los resultados de esa sección; si no es posible filtrarlo, indícalo explícitamente.\n"
    "- Si hay más de 10 resultados, agrupa por departamento o rango normativo.\n"
    "- Para sumarios BOE: muestra primero las Secciones I y II (disposiciones y personal), "
    "luego el resto si el usuario lo solicita.\n"
    "- Termina siempre con la fuente entre paréntesis: (Fuente: BOE sumario YYYY-MM-DD) "
    "o el nombre del documento PDF según corresponda."
)


@dataclass
class ContextDefinition:
    key: str
    label: str
    description: str
    color: str
    content: str


# ---------------------------------------------------------------------------
# File readers
# ---------------------------------------------------------------------------


def _read_docx(path: Path) -> str:
    try:
        from docx import Document  # type: ignore[import]

        doc = Document(str(path))
        parts = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        parts.append(cell.text.strip())
        return "\n".join(parts)
    except ImportError:
        print(
            "[WARN] python-docx no instalado - contenido .docx no disponible. Ejecuta: uv sync"
        )
        return ""
    except Exception as exc:
        print(f"[WARN] Error leyendo {path.name}: {exc}")
        return ""


def _load_team_activities_txt() -> str:
    path = CONTEXTS_DIR / "team_activities.txt"
    if path.exists():
        return path.read_text(encoding="utf-8")
    return ""


def _parse_team_sections(txt: str) -> dict[str, str]:
    """
    Parse the team activities txt into numbered sections (1-8) plus a 'limits' entry.
    Returns a dict: {"1": "section 1 text", ..., "8": "...", "limits": "..."}
    """
    # Match lines that start sections: "       1) EFICIENCIA..."
    pattern = re.compile(r"^\s{2,}(\d+)\)\s+", re.MULTILINE)
    matches = list(pattern.finditer(txt))

    sections: dict[str, str] = {}
    for i, match in enumerate(matches):
        num = match.group(1)
        start = match.start()
        end = matches[i + 1].start() if i + 1 < len(matches) else len(txt)
        sections[num] = txt[start:end].strip()

    # Extract LÍMITES / ANTI-AJUSTE section (appears after numbered sections)
    limits_match = re.search(r"LÍMITES\s*/\s*ANTI-AJUSTE", txt)
    if limits_match:
        sections["limits"] = txt[limits_match.start() :].strip()

    return sections


# ---------------------------------------------------------------------------
# Context content builders
# ---------------------------------------------------------------------------


def _build_context(section_txt: str, extra: str = "") -> str:
    """Combine a section text with the limits block and any extra content."""
    _txt = _load_team_activities_txt()
    sections = _parse_team_sections(_txt) if _txt else {}
    limits = sections.get("limits", "")

    parts = [section_txt]
    if limits:
        parts.append(limits)
    if extra:
        parts.append(extra)
    return "\n\n".join(p for p in parts if p)


def _load_all_contexts() -> dict[str, ContextDefinition]:
    """Build all team context definitions from the team activities txt (and SGE docx)."""
    txt = _load_team_activities_txt()
    sections = _parse_team_sections(txt) if txt else {}
    limits = sections.get("limits", "")

    def _content(*nums: str, extra: str = "") -> str:
        parts = [sections[n] for n in nums if n in sections]
        if limits:
            parts.append(limits)
        if extra:
            parts.append(extra)
        return "\n\n".join(p for p in parts if p)

    # SGE docx extra prompt
    docx_extra = ""
    docx_path = CONTEXTS_DIR / "Prompt BOE.docx"
    if docx_path.exists():
        docx_extra = _read_docx(docx_path)

    return {
        "sge": ContextDefinition(
            key="sge",
            label="SGE — Sistemas de Gestión Energética",
            description="Eficiencia energética, ISO 50001, auditorías, descarbonización, comunidades energéticas.",
            color="teal",
            content=_content("1", extra=docx_extra),
        ),
        "grr": ContextDefinition(
            key="grr",
            label="GRR — Generación y Recursos Renovables",
            description="Energía eólica y solar FV, curvas de potencia, análisis de recurso, hibridación, O&M.",
            color="emerald",
            content=_content("2"),
        ),
        "ric": ContextDefinition(
            key="ric",
            label="RIC — Redes Inteligentes y Control",
            description="Subestaciones, IEC 61850, protecciones, microrredes, EMS, ciberseguridad OT.",
            color="blue",
            content=_content("3"),
        ),
        "mev": ContextDefinition(
            key="mev",
            label="MEV — Movilidad Eléctrica y Vehículos",
            description="Recarga VE, V2G, OCPP, infraestructura de carga, PMUS, microsimulaciones de tráfico.",
            color="violet",
            content=_content("4"),
        ),
        "iad": ContextDefinition(
            key="iad",
            label="IAD — IA y Digitalización",
            description="Gemelos digitales, dataspaces, visión artificial, optimización, IA aplicada a energía e industria.",
            color="orange",
            content=_content("5"),
        ),
        "hid": ContextDefinition(
            key="hid",
            label="HID — Hidrógeno y Descarbonización Industrial",
            description="Hidrógeno verde, Power-to-X, electrolizadores, valles de hidrógeno, electrificación industrial.",
            color="sky",
            content=_content("6"),
        ),
        "rec": ContextDefinition(
            key="rec",
            label="REC — Reciclado, Economía Circular y Sostenibilidad",
            description="ACV, ecodiseño, reciclado de baterías y paneles FV, CSRD, ISO 14000, bioeconomía.",
            color="lime",
            content=_content("7"),
        ),
        "tis": ContextDefinition(
            key="tis",
            label="TIS — Tecnologías Industriales y Simulación",
            description="CFD, FEM, simulación industrial, inducción, modelado de procesos termoquímicos.",
            color="slate",
            content=_content("8"),
        ),
    }


# Registry is populated once at import time
CONTEXT_REGISTRY: dict[str, ContextDefinition] = _load_all_contexts()


# ---------------------------------------------------------------------------
# Public API
# ---------------------------------------------------------------------------


def get_team_context_text(active_context_keys: list[str]) -> str:
    """Return only the team context blocks as a single string (without BASE_SYSTEM_PROMPT).

    Used by llm_handler to build a dedicated cache_control block for the team context,
    separate from the identity and format blocks.
    Returns empty string when no valid team context is active.
    """
    blocks: list[str] = []
    for key in active_context_keys:
        ctx = CONTEXT_REGISTRY.get(key.lower())
        if ctx and ctx.content:
            blocks.append(
                f"--- CONTEXTO DE EQUIPO: {ctx.label} ---\n"
                f"Eres Tambora asistiendo al equipo {ctx.key.upper()}. "
                f"Cuando respondas consultas sobre el BOE, ten en cuenta las siguientes "
                f"capacidades y areas de trabajo del equipo para enriquecer el contexto "
                f"y priorizar la relevancia de tus respuestas:\n\n"
                f"{ctx.content}\n\n"
                f"--- FIN CONTEXTO DE EQUIPO {ctx.key.upper()} ---"
            )
    return "\n\n".join(blocks)


def get_available_contexts() -> list[dict]:
    return [
        {
            "key": ctx.key,
            "label": ctx.label,
            "description": ctx.description,
            "color": ctx.color,
        }
        for ctx in CONTEXT_REGISTRY.values()
    ]


def build_system_prompt(active_context_keys: list[str]) -> str:
    """Return BASE_SYSTEM_PROMPT, optionally prepended with team context blocks."""
    if not active_context_keys:
        return BASE_SYSTEM_PROMPT

    blocks: list[str] = []
    for key in active_context_keys:
        ctx = CONTEXT_REGISTRY.get(key.lower())
        if ctx and ctx.content:
            blocks.append(
                f"--- CONTEXTO DE EQUIPO: {ctx.label} ---\n"
                f"Eres Tambora asistiendo al equipo {ctx.key.upper()}. "
                f"Cuando respondas consultas sobre el BOE, ten en cuenta las siguientes "
                f"capacidades y areas de trabajo del equipo para enriquecer el contexto "
                f"y priorizar la relevancia de tus respuestas:\n\n"
                f"{ctx.content}\n\n"
                f"--- FIN CONTEXTO DE EQUIPO {ctx.key.upper()} ---"
            )

    if not blocks:
        return BASE_SYSTEM_PROMPT

    return "\n\n".join(blocks) + "\n\n" + BASE_SYSTEM_PROMPT
